"""本機 Word 文件及沿既有 house CSS 的單檔文章建置。"""
from pathlib import Path
import re, html, json
import markdown
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'monjaro/deliverables'; QA=Path(__file__).parent
text=(OUT/'mounjaro_paper_explainer.md').read_text(encoding='utf-8')
assert '??' not in text
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.75)
sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    style=doc.styles[name];style.font.name='Microsoft JhengHei';style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei')
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.2
doc.styles['Title'].font.size=Pt(22);doc.styles['Heading 1'].font.size=Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(15)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
doc.core_properties.author='[TO CONFIRM]';doc.core_properties.title='Tirzepatide 與 semaglutide 的臨床證據解讀'
lines=text.splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            vals=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(set(x)<=set('- :') for x in vals):rows.append(vals)
            i+=1
        table=doc.add_table(rows=1,cols=len(rows[0]));table.autofit=False
        for j,v in enumerate(rows[0]):table.rows[0].cells[j].text=v
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for row in rows[1:]:
            cells=table.add_row().cells
            for j,v in enumerate(row):cells[j].text=v
        for rn,row in enumerate(table.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(6);p.paragraph_format.space_before=Pt(5)
                    for run in p.runs:run.font.size=Pt(9);run.bold=rn==0
                if rn==0:
                    shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E7ECEF');cell._tc.get_or_add_tcPr().append(shading)
        continue
    if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
    elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
    else:
        p=doc.add_paragraph(line.replace('**',''))
        if re.match(r'S\d\d ',line):
            for run in p.runs:run.font.size=Pt(9)
    i+=1
doc.save(OUT/'mounjaro_paper_explainer.docx')
house=(ROOT/'monjaro/sources/house_style/kidney-cognition-lab/blog/vbm-ei-alzheimers.html').read_text(encoding='utf-8')
css=re.search(r'<style>(.*?)</style>',house,re.S).group(1)
body=markdown.markdown('\n'.join(lines[6:]),extensions=['tables'])
body=body.replace('<table>','<div class="tbl-wrap"><table>').replace('</table>','</table></div>')
body=re.sub(r'(https://[^\s<]+)',lambda m:f'<a href="{html.escape(m.group(1),quote=True)}">{html.escape(m.group(1))}</a>',body)
extra='''body{font-family:"Microsoft JhengHei","Noto Sans TC",sans-serif}svg{display:block;width:100%;height:auto}button{min-height:44px}button:focus-visible,a:focus-visible{outline:3px solid var(--gold);outline-offset:3px}.dashboard{padding:20px}.dashboard h2{margin-top:0}.dashboard canvas{width:100%}#context-text{color:var(--text);margin:15px 0}td,th{min-width:150px}p,a{overflow-wrap:anywhere}@media(max-width:640px){.dashboard{padding:14px}.article-title{font-size:1.7rem}.preset-grid{grid-template-columns:1fr 1fr}}@media(prefers-reduced-motion:reduce){*{scroll-behavior:auto!important;animation:none!important;transition:none!important}}'''
hero='''<figure class="hero-figure"><svg viewBox="0 0 800 200" role="img" aria-label="從臨床問題對齊族群介入與終點"><rect width="800" height="200" rx="20" fill="#111d3c"/><g fill="#e8f4f8" font-family="Microsoft JhengHei,sans-serif" text-anchor="middle"><text x="400" y="45" font-size="23">先對齊研究問題 再比較藥物</text><text x="130" y="112" font-size="22" fill="#4dd9cf">誰接受治療</text><text x="400" y="112" font-size="22" fill="#a78bfa">用什麼跟誰比</text><text x="670" y="112" font-size="22" fill="#f4a261">量到什麼結果</text><text x="400" y="163" font-size="18">分子 · 品牌 · 劑量 · 族群 · 追蹤 · 終點 · 地區</text></g></svg><figcaption>機制、體重、症狀與臨床事件分開閱讀。</figcaption></figure>'''
interactive=(QA/'interactive.html').read_text(encoding='utf-8')
script=(QA/'interactive.js').read_text(encoding='utf-8')
page=f'''<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="robots" content="noindex,nofollow"><meta name="description" content="釐清 tirzepatide 與 semaglutide 在控糖、肥胖、HFpEF 與 CKD 的證據邊界。"><title>Tirzepatide 與 semaglutide 的臨床證據解讀</title><style>{css}\n{extra}</style></head><body><nav class="top-bar"><div class="top-bar-inner"><a href="#article" class="back-link">閱讀全文</a><span class="top-bar-title">醫療專業教育 · 本機草稿</span></div></nav><main class="article" id="article"><header class="article-header"><div class="article-meta">2026-10-31 · 糖尿病 肥胖 心腎證據</div><h1 class="article-title">Tirzepatide 與 semaglutide<br>先對齊研究問題 再比較藥物</h1><p class="article-subtitle">從控糖、減重、肥胖相關心衰竭到腎臟風險，學會把數字放回正確的族群、劑量與終點。</p><div class="article-author"><div class="author-info">作者：[TO CONFIRM]　機構：[TO CONFIRM]　COI：[TO CONFIRM]</div></div></header>{hero}{interactive}{body}<footer class="article-footer">供醫療專業教育，不作個別病人診療或處方。台灣完整仿單、給付與器型需另核實。此為本機草稿；尚未發布。</footer></main><script>{script}</script></body></html>'''
(OUT/'homepage_blog_draft.html').write_text(page,encoding='utf-8')
print(json.dumps({'docx':str(OUT/'mounjaro_paper_explainer.docx'),'blog':str(OUT/'homepage_blog_draft.html'),'characters':len(text)},ensure_ascii=False))
